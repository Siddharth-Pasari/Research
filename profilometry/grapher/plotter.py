import tkinter as tk
from tkinter import filedialog

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import openpyxl as px
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.trendline import Trendline
from openpyxl.styles import Font

def update_selected_value(choice, sections, selected_value):
    matching_sections = [section for section in sections if section[1] == choice]
    selected_value[0] = matching_sections[0] if matching_sections else (None, None)

def get_section_differences(data_nl, selected_value):
    difference_list = [[]]
    val = selected_value[0][0]

    while data_nl[3][val] != "Difference" and val + 1 < len(data_nl[3]):  # This would indicate the section ended
        if pd.isna(data_nl[3][val]):
            difference_list.append([])
        else:
            if data_nl[3][val] in ('NoVal', 'No Value'):
                difference_list[-1].append(0)
            else:
                difference_list[-1].append(data_nl[3][val])
        val += 1

    return difference_list[:-2]  # Remove the two blank entries


def generate_scatter_plot(difference_list):
    ftnum = len(difference_list[0])  # Number of features per repetition
    sums = [0] * ftnum
    counts = [0] * ftnum

    for sublist in difference_list:
        for index in range(len(sublist)):
            sums[index] += sublist[index]
            counts[index] += 1

    y_values = [sums[i] / counts[i] if counts[i] > 0 else float('nan') for i in range(ftnum)]
    x_values = np.arange(len(y_values))

    fig, axs = plt.subplots()
    axs.scatter(x_values, y_values)

    if len(x_values) > 1:
        coefficients = np.polyfit(x_values, y_values, 1)
        trend_line = np.poly1d(coefficients)
        axs.plot(x_values, trend_line(x_values), color='red', linewidth=1)

    axs.set_xlabel('Feature Number')
    axs.set_ylabel('Height (µm)')
    axs.set_title('Scatter Plot of Logged Features')
    plt.show()


def add_graph_to_excel(path, difference_list, selected_value):
    workbook = px.load_workbook(path)
    graph_sheet_name = f"Graph - {selected_value[0][1]}"
    worksheet = workbook.create_sheet(graph_sheet_name)
    
    ftnum = len(difference_list[0])
    sums = [0] * ftnum
    counts = [0] * ftnum
    
    for sublist in difference_list:
        for index in range(len(sublist)):
            sums[index] += sublist[index]
            counts[index] += 1
    y_values = [sums[i] / counts[i] if counts[i] > 0 else float('nan') for i in range(ftnum)]
    x_values = np.arange(len(y_values))
    
    for row, (x, y) in enumerate(zip(x_values, y_values), start=1):
        worksheet[f'A{row}'] = x
        worksheet[f'B{row}'] = y
    
    coefficients = np.polyfit(x_values, y_values, 1)
    trendline_values = np.polyval(coefficients, x_values)
    
    for row, value in enumerate(trendline_values, start=1):
        worksheet[f'C{row}'] = value
    
    chart = ScatterChart()
    chart.title = selected_value[0][1]
    chart.x_axis.title = 'Feature Number'
    chart.y_axis.title = 'Height (µm)'
    
    x_ref = Reference(worksheet, min_col=1, min_row=1, max_row=len(x_values))
    y_ref = Reference(worksheet, min_col=2, min_row=1, max_row=len(y_values))
    series = Series(y_ref, x_ref, title="Scatter")
    series.marker.symbol = 'circle'
    series.marker.size = 7
    series.graphicalProperties.line.noFill = True
    chart.series.append(series)
    
    trendline_series = Series(Reference(worksheet, min_col=3, min_row=1, max_row=len(trendline_values)), x_ref, title="Trendline")
    trendline_series.graphicalProperties.line.solidFill = "FF0000"
    trendline_series.graphicalProperties.line.width = 20000
    chart.series.append(trendline_series)
    
    chart.x_axis.num_font = Font(bold=True)
    chart.y_axis.num_font = Font(bold=True)
    chart.x_axis.majorGridlines = None
    chart.y_axis.majorGridlines = None
    chart.legend = None
    
    worksheet.add_chart(chart, "E5")
    workbook.save(path)


def main():
    root = tk.Tk()
    root.withdraw()
    
    path = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx")])
    data_df = pd.read_excel(path)
    data_nl = data_df.values.T.tolist()
    sections = [(i + 1, data_nl[0][i - 1]) for i, value in enumerate(data_nl[0]) if value == "Num"]
    
    root.deiconify()
    selected_value = [sections[0]]
    options = [section[1] for section in sections]
    
    dropdown_menu = tk.OptionMenu(root, tk.StringVar(root, value=options[0]), *options, command=lambda choice: update_selected_value(choice, sections, selected_value))
    dropdown_menu.pack()
    
    tk.Button(root, text="Generate Scatter Plot", command=lambda: generate_scatter_plot(get_section_differences(data_nl, selected_value))).pack()
    tk.Button(root, text="Add Graph to Excel", command=lambda: add_graph_to_excel(path, get_section_differences(data_nl, selected_value), selected_value)).pack()
    
    root.mainloop()


if __name__ == "__main__":
    main()